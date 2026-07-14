# -*- coding: utf-8 -*-
"""
四年级暑假作业计划表生成器(2026年暑假 7月13日—8月31日,共8周)
依据四张作业单:语文暑假作业、数学暑假作业、英语暑假作业、体育家庭任务单(北京林业大学附属小学)
生成:2 个总览页(竖版)+ 8 个周计划页(横向),每页均可打印为一张 A4。
"""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.pagebreak import Break
from openpyxl.utils import get_column_letter

# ---------- 颜色 ----------
NAVY = "1F4E79"          # 标题
HEADER_BLUE = "4472C4"   # 表头(工作日)
WEEKEND_HD = "C55A11"    # 表头(周末)——深橙,保证白字对比度,黑白打印也清晰
CHN_D, CHN_L = "C00000", "FDEAEA"        # 语文
MATH_D, MATH_L = "2E75B6", "DDEBF7"      # 数学
ENG_D, ENG_L = "548235", "E2EFDA"        # 英语
PE_D, PE_L = "C55A11", "FCE4D6"          # 体育
PRAC_D, PRAC_L = "7030A0", "EDE4F5"      # 数学实践
GRAY = "F2F2F2"          # 休息格
NOTE_BG = "FFF2CC"       # 底部提示
TAB_COLORS = ["E74C3C", "E67E22", "F1C40F", "2ECC71", "1ABC9C", "3498DB", "9B59B6", "E91E63"]

FONT = "微软雅黑"
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

WD_NAMES = ["周一", "周二", "周三", "周四", "周五", "周六", "周日"]
START = dt.date(2026, 7, 13)  # 第1周周一


def fill(color):
    return PatternFill("solid", fgColor=color)


def set_cell(ws, row, col, value, *, font=None, fl=None, align=None, border=True):
    c = ws.cell(row=row, column=col, value=value)
    if font:
        c.font = font
    if fl:
        c.fill = fill(fl)
    if align:
        c.alignment = align
    if border:
        c.border = BORDER
    return c


CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)  # 缩进,文字不贴边框


def setup_page(ws):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9  # A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.page_margins.top = ws.page_margins.bottom = 0.35
    ws.print_options.horizontalCentered = True


def scale_rows(ws, last_row, target=780):
    """按比例放大行高,让表格纵向铺满 A4(fitToWidth 缩放后仍有留白时使用)。"""
    total = sum(ws.row_dimensions[r].height or 15 for r in range(1, last_row + 1))
    factor = min(max(target / total, 1.0), 2.0)
    for r in range(1, last_row + 1):
        ws.row_dimensions[r].height = round((ws.row_dimensions[r].height or 15) * factor, 1)


def day_dates(week):  # week: 1..8
    monday = START + dt.timedelta(days=7 * (week - 1))
    return [monday + dt.timedelta(days=i) for i in range(7)]


# =====================================================================
# 周计划页(第1~7周)
# =====================================================================
def build_week(wb, week):
    dates = day_dates(week)
    ws = wb.create_sheet(f"第{week}周({dates[0].month}.{dates[0].day}-{dates[6].month}.{dates[6].day})")
    ws.sheet_properties.tabColor = TAB_COLORS[week - 1]
    setup_page(ws)

    widths = {"A": 6.5, "B": 33, "J": 22}
    for col in "CDEFGHI":
        widths[col] = 14.5
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    # 标题
    ws.merge_cells("A1:J1")
    set_cell(ws, 1, 1,
             f"四年级暑假每日安排 · 第{week}周({dates[0].month}月{dates[0].day}日—{dates[6].month}月{dates[6].day}日)"
             f"    完成一项在□里打✓",
             font=Font(name=FONT, size=15, bold=True, color="FFFFFF"), fl=NAVY, align=CENTER)
    ws.row_dimensions[1].height = 30

    # 表头
    set_cell(ws, 2, 1, "科目", font=Font(name=FONT, size=11, bold=True, color="FFFFFF"), fl=HEADER_BLUE, align=CENTER)
    set_cell(ws, 2, 2, "任务内容", font=Font(name=FONT, size=11, bold=True, color="FFFFFF"), fl=HEADER_BLUE, align=CENTER)
    for i, d in enumerate(dates):
        hd = WEEKEND_HD if i >= 5 else HEADER_BLUE
        set_cell(ws, 2, 3 + i, f"{WD_NAMES[i]}\n{d.month}月{d.day}日",
                 font=Font(name=FONT, size=11, bold=True, color="FFFFFF"), fl=hd, align=CENTER)
    set_cell(ws, 2, 10, "要求 / 备注", font=Font(name=FONT, size=11, bold=True, color="FFFFFF"),
             fl=HEADER_BLUE, align=CENTER)
    ws.row_dimensions[2].height = 30

    row = 3
    section_rows = {}  # subject -> (start,end)

    def task_row(subject, task, cells, note, light, *, height=24, merge_days=False, left_days=False):
        nonlocal row
        set_cell(ws, row, 2, task, font=Font(name=FONT, size=10.5, bold=False), fl=light, align=LEFT)
        if merge_days:
            ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=9)
            set_cell(ws, row, 3, cells, font=Font(name=FONT, size=10.5), fl=light, align=CENTER)
            for cc in range(4, 10):
                set_cell(ws, row, cc, None, fl=light)
        else:
            for i in range(7):
                v = cells.get(i, "—")
                bg = light if v not in ("—", "休息 · 自由活动") else GRAY
                if left_days and v != "—":
                    fnt, al = Font(name=FONT, size=10.5), LEFT   # 多行清单左对齐更规整
                elif isinstance(v, str) and "□" in v:
                    fnt, al = Font(name=FONT, size=12), CENTER   # 打钩方框放大
                else:
                    fnt, al = Font(name=FONT, size=10.5), CENTER
                set_cell(ws, row, 3 + i, v, font=fnt, fl=bg, align=al)
        set_cell(ws, row, 10, note, font=Font(name=FONT, size=10), fl=light, align=LEFT)
        ws.row_dimensions[row].height = height
        section_rows.setdefault(subject, [row, row])[1] = row
        if subject not in section_rows or section_rows[subject][0] > row:
            section_rows[subject][0] = row
        row += 1

    WD = {i: "□" for i in range(5)}  # 周一~周五打卡

    # ---------------- 语文 ----------------
    task_row("语文", "必做① 抄写1课《词语表》词语,每词2遍(写在生字本)", WD,
             "语文书P142-144;书写规范、工整、美观", CHN_L)
    task_row("语文", "必做② 背诵《208篇》初中部分(全假期不少于20篇)",
             {0: "□ 1篇", 2: "□ 1篇", 4: "□ 1篇"},
             "每周一、三、五各1篇,第1~7周共21篇;声音洪亮、字音准确", CHN_L)
    if week == 2:
        task_row("语文", "必做·习作①《写写我的心爱之物》(作文纸)",
                 {4: "□ 完成习作"}, "是什么、什么样、怎么得到、为什么成为心爱之物,表达喜爱之情", CHN_L, height=28)
    if week == 5:
        task_row("语文", "必做·习作②《慧眼看世界,妙笔游美景》(作文纸)",
                 {4: "□ 完成习作"}, "写假期一处美景,按一定顺序写出动态变化;有条件可录旅行视频(横屏、声音清晰、≥1分钟)", CHN_L, height=28)
    if week == 6:
        task_row("语文", "选做:“心灵奇旅”影片人物推荐卡",
                 {4: "□ 制作推荐卡"}, "与家人看一场电影;含人物名称/特点/画像/感受;注明电影名称,用典型事例介绍人物", CHN_L, height=28)
    task_row("语文", "选做:每日阅读感兴趣的书10分钟+《我的阅读记录卡》",
             {i: "□ 10分钟" for i in range(5)},
             "推荐《中外民间故事》及名著、名家散文;开学“好书我推荐”2分钟(可做PPT)", CHN_L)
    task_row("语文", "选做:练字一页字帖+《练字记录表》自评", WD,
             "个别同学选做,书写美观者免做;开学后装订成册上交", CHN_L)
    chn_end = row - 1
    math_start = row

    # ---------------- 数学 ----------------
    task_row("数学", "口算 15~20 道(题目自备)", WD, "每日基础练习", MATH_L)
    task_row("数学", "计算 4 道:三位数÷两位数 2道 + 小数乘法 2道", WD, "每日基础练习", MATH_L)
    if week <= 6:
        paper = {5: f"□ 第{week}份(必做)", 6: "□ 加练1份(选做)"}
    else:  # 第7周补齐8周的量
        paper = {5: "□ 第7份(必做)", 6: "□ 第8份(必做)"}
    task_row("数学", "综合练习卷(每周1~2份,按8周计)", paper,
             "三选一:①数练活页卷 ②《5.3全优卷》剩余试卷 ③自备练习卷", MATH_L, height=28)
    if week == 4:
        cells = {i: "□ 记当天气温" for i in range(6)}
        cells[6] = "□ 记气温+绘统计图,提数学问题并解答"
        task_row("数学", "实践作业(原单第③项):统计一周(7天)气温", cells,
                 "A4纸完成,配表格/照片(实践4选2之一)", PRAC_L, height=32)
    if week == 6:
        cells = {6: "□ 设计一幅美丽的密铺图案"}
        task_row("数学", "实践作业(原单第④项):设计密铺图案", cells,
                 "一种图形或几种图形组合密铺,A4纸完成(实践4选2之二)", PRAC_L, height=28)
    task_row("数学", "选做:《教材全解》拓展题 / 与家长商定提高内容",
             "学有余力时安排,不做硬性打卡", "自行安排", MATH_L, merge_days=True, height=20)

    math_end = row - 1

    # ---------------- 英语 ----------------
    eng_start = row
    MTH = {i: "□ 5分钟" for i in range(4)}  # 周一~周四
    if week <= 6:
        task_row("英语", "必做① 每天大声朗读三下课本(每次5分钟)", MTH,
                 f"全假期共24次;第1~6周每周4次(本周为第{week*4-3}~{week*4}次)", ENG_L)
        task_row("英语", "必做② 抄写四上&四下单词/词组并造句(每次10分钟)",
                 {i: "□ 10分钟" for i in range(4)}, "全假期共24次;第1~6周每周4次", ENG_L)
    if week <= 5:
        task_row("英语", "必做③ 完成英语阅读题3篇(每次15分钟)",
                 {0: "□ 3篇", 2: "□ 3篇", 4: "□ 3篇"}, "全假期共15次;第1~5周每周一、三、五", ENG_L)
    if week in (5, 6):
        task_row("英语", "选做① 学唱一首英文歌曲(歌曲自选,每次5分钟)",
                 {4: "□ 5分钟"}, "共2次(第5、6周各1次)", ENG_L)
    if week == 7:
        task_row("英语", "必做④ 学科实践:选3~4首喜爱的古诗,完成A3海报(每次15分钟)",
                 {0: "□ 选古诗\n构思", 1: "□ 设计\n版面草稿", 2: "□ 抄写古诗", 3: "□ 配图上色", 4: "□ 完善\n落款完成"},
                 "共5次;古诗最好是自己学过的中文古诗", ENG_L, height=34)
        task_row("英语", "选做② 英语试听练习:英文动画电影 / BBC记录片(每次≤10分钟)",
                 {2: "□ ≤10分钟", 4: "□ ≤10分钟"}, "共2次,内容自选", ENG_L)
    eng_end = row - 1

    # ---------------- 体育 ----------------
    pe_start = row
    pe_daily = {
        0: "□ 热身\n□ 1分钟计时跳绳×4组\n(记成绩)\n□ 坐位体前屈拉伸1分钟",
        1: "□ 热身\n□ 跳绳200个/组×4组\n□ 1分钟仰卧起坐×2组\n(记成绩)\n□ 坐位体前屈1分钟",
        2: "□ 热身\n□ 1分钟计时跳绳×4组\n(记成绩)\n□ 吹气球练肺活量×3组\n□ 放松拉伸",
        3: "□ 热身\n□ 3分钟耐力跳绳\n(自选音乐,记总个数)\n□ 放松拉伸",
        4: "□ 热身\n□ 耐力跑2公里\n(避开中午炎热时段,\n注意防暑降温,并记录)\n□ 放松拉伸",
        5: "休息 · 自由活动",
        6: "休息 · 自由活动",
    }
    task_row("体育", "每日锻炼(先热身 → 练习 → 放松拉伸)", pe_daily,
             "目标:跳绳200个/分钟;仰卧起坐49个/分钟;体前屈膝盖伸直、手指过脚尖;有条件可加练加速跑",
             PE_L, height=96, left_days=True)
    task_row("体育", "成绩记录(把当天成绩写进括号)",
             {i: "(        )" for i in range(5)}, "开学交任务单,学校按完成情况评优体奖章", PE_L, height=26)
    # 家长签字行
    set_cell(ws, row, 2, "家长点评签字(每周一次)", font=Font(name=FONT, size=10.5, bold=True), fl=PE_L, align=LEFT)
    ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=9)
    set_cell(ws, row, 3, "", fl="FFFFFF")
    for cc in range(4, 10):
        set_cell(ws, row, cc, None, fl="FFFFFF")
    set_cell(ws, row, 10, "对应任务单“家长点评签字”栏", font=Font(name=FONT, size=10), fl=PE_L, align=LEFT)
    ws.row_dimensions[row].height = 26
    pe_end = row
    row += 1

    # 科目竖排色块
    for (subject, dark, r1, r2) in [("语文", CHN_D, 3, chn_end),
                                    ("数学", MATH_D, math_start, math_end),
                                    ("英语", ENG_D, eng_start, eng_end),
                                    ("体育", PE_D, pe_start, pe_end)]:
        ws.merge_cells(start_row=r1, start_column=1, end_row=r2, end_column=1)
        set_cell(ws, r1, 1, subject, font=Font(name=FONT, size=13, bold=True, color="FFFFFF"),
                 fl=dark, align=CENTER)
        for r in range(r1 + 1, r2 + 1):
            set_cell(ws, r, 1, None, fl=dark)

    # 底部提示
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
    tips = ("★ 语文、英语作业不含周六日和法定节假日 · 体育锻炼注意防暑降温 · "
            "数学实践作业4选2(本表建议做③气温统计、④密铺图案,也可换成①数学游戏、②小区调查)")
    if week == 7:
        tips = ("★ 本周是体育任务单最后一周,周末把“注明”栏填好:跳绳最好成绩、跑步公里数及用时、仰卧起坐最好成绩 · "
                "语文背诵本周累计满21篇 · 英语必做已全部排完,下周一(8月31日)填自评单")
    set_cell(ws, row, 1, tips, font=Font(name=FONT, size=10.5, bold=True, color="7F6000"),
             fl=NOTE_BG, align=LEFT)
    for cc in range(2, 11):
        set_cell(ws, row, cc, None, fl=NOTE_BG)
    ws.row_dimensions[row].height = 26
    scale_rows(ws, row)


# =====================================================================
# 第8周 收尾页
# =====================================================================
def build_week8(wb):
    ws = wb.create_sheet("第8周(8.31收尾)")
    ws.sheet_properties.tabColor = TAB_COLORS[7]
    setup_page(ws)
    for col, w in {"A": 6.5, "B": 46, "C": 30, "D": 46}.items():
        ws.column_dimensions[col].width = w

    ws.merge_cells("A1:D1")
    set_cell(ws, 1, 1, "四年级暑假 · 第8周 收尾冲刺(8月31日 周一)     9月1日开学!",
             font=Font(name=FONT, size=16, bold=True, color="FFFFFF"), fl=NAVY, align=CENTER)
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:B2")
    set_cell(ws, 2, 1, "8月31日(周一)当日安排", font=Font(name=FONT, size=12, bold=True, color="FFFFFF"),
             fl=HEADER_BLUE, align=CENTER)
    set_cell(ws, 2, 2, None, fl=HEADER_BLUE)
    ws.merge_cells("C2:D2")
    set_cell(ws, 2, 3, "开学上交材料清单(装进书包!)", font=Font(name=FONT, size=12, bold=True, color="FFFFFF"),
             fl="C00000", align=CENTER)
    set_cell(ws, 2, 4, None, fl="C00000")
    ws.row_dimensions[2].height = 26

    today_tasks = [
        ("语文", "□ 抄写最后1课词语;核对背诵是否满20篇(计划已排21篇)", CHN_L, CHN_D),
        ("语文", "□ 准备开学实践“好书我推荐”2分钟发言(可制作PPT);检查两篇必做习作、阅读记录卡、练字册(选做)", CHN_L, CHN_D),
        ("数学", "□ 口算15~20道 + 计算4道(最后一组)", MATH_L, MATH_D),
        ("数学", "□ 清点综合练习卷:必做共8份是否完成、订正", MATH_L, MATH_D),
        ("数学", "□ 检查2项实践作业(A4纸、配照片或表格)", MATH_L, MATH_D),
        ("英语", "□ 填写《暑假英语学习自评单》:班级、姓名,勾选2道“是/否”,写下学期英语课的想法和建议", ENG_L, ENG_D),
        ("英语", "□ 检查A3古诗海报、抄写造句和阅读题是否齐全", ENG_L, ENG_D),
        ("体育", "□ 自由锻炼30分钟(跳绳/慢跑,注意防暑)", PE_L, PE_D),
        ("体育", "□ 体育任务单“注明”栏:填跳绳最好成绩、跑步公里数及用时、仰卧起坐最好成绩;7周家长点评签字补齐", PE_L, PE_D),
        ("整理", "□ 收拾书包文具,调整作息,早睡早起迎接开学", NOTE_BG, "BF8F00"),
    ]
    submit_list = [
        ("语文", "□ 生字本(每日抄写词语)+ 两篇必做习作(作文纸)", CHN_L, CHN_D),
        ("语文", "□ 练字册(选做者装订成册)+ 阅读记录卡、影片推荐卡(选做)+“好书我推荐”PPT", CHN_L, CHN_D),
        ("数学", "□ 综合练习卷 8 份(活页卷/《5.3全优卷》/自备)", MATH_L, MATH_D),
        ("数学", "□ 口算、计算练习记录", MATH_L, MATH_D),
        ("数学", "□ 实践性作业 2 项(A4纸,自设计格式,配照片或表格)", MATH_L, MATH_D),
        ("英语", "□ 《暑假英语学习自评单》(开学报到第一天上交)", ENG_L, ENG_D),
        ("英语", "□ A3古诗海报作品(学科实践活动)", ENG_L, ENG_D),
        ("英语", "□ 单词抄写造句本、英语阅读题", ENG_L, ENG_D),
        ("体育", "□ 体育家庭任务单(成绩记录+家长点评签字+最好成绩注明)", PE_L, PE_D),
        ("其他", "□ 学校/班级另行通知的材料", NOTE_BG, "BF8F00"),
    ]
    r = 3
    for i in range(max(len(today_tasks), len(submit_list))):
        if i < len(today_tasks):
            sub, txt, light, dark = today_tasks[i]
            set_cell(ws, r, 1, sub, font=Font(name=FONT, size=11, bold=True, color="FFFFFF"), fl=dark, align=CENTER)
            set_cell(ws, r, 2, txt, font=Font(name=FONT, size=11.5), fl=light, align=LEFT)
        if i < len(submit_list):
            sub, txt, light, dark = submit_list[i]
            set_cell(ws, r, 3, f"{sub} | {txt}", font=Font(name=FONT, size=11.5), fl=light, align=LEFT)
            set_cell(ws, r, 4, "", fl=light)
        ws.row_dimensions[r].height = 34
        r += 1
    # 合并上交清单两列
    for rr in range(3, r):
        ws.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=4)

    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    set_cell(ws, r, 1, "★ 8月29日(周六)、30日(周日)属于第7周页:综合练习卷第7、8份安排在这两天。"
                        "如有未完成项目,利用这三天全部补齐!",
             font=Font(name=FONT, size=10.5, bold=True, color="7F6000"), fl=NOTE_BG, align=LEFT)
    ws.row_dimensions[r].height = 28
    scale_rows(ws, r)


# =====================================================================
# 总览页(拆为两个工作表,各一页A4竖版,保证字号可读)
# =====================================================================
def build_overview(wb):
    state = {"ws": None, "r": 2}

    def start_sheet(sheet_name, title_text, first=False):
        if first:
            ws = wb.active
            ws.title = sheet_name
        else:
            ws = wb.create_sheet(sheet_name)
        ws.sheet_properties.tabColor = NAVY
        setup_page(ws)
        ws.page_setup.orientation = "portrait"  # 总览内容纵向长,竖版利用率高
        for col, w in {"A": 8, "B": 48, "C": 30, "D": 30, "E": 30}.items():
            ws.column_dimensions[col].width = w
        ws.merge_cells("A1:E1")
        set_cell(ws, 1, 1, title_text,
                 font=Font(name=FONT, size=15, bold=True, color="FFFFFF"), fl=NAVY, align=CENTER)
        ws.row_dimensions[1].height = 32
        state["ws"] = ws
        state["r"] = 2
        return ws

    def section(title, dark):
        ws, r = state["ws"], state["r"]
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        set_cell(ws, r, 1, title, font=Font(name=FONT, size=12.5, bold=True, color="FFFFFF"), fl=dark, align=LEFT)
        for cc in range(2, 6):
            set_cell(ws, r, cc, None, fl=dark)
        ws.row_dimensions[r].height = 24
        state["r"] = r + 1

    def head_row(light):
        ws, r = state["ws"], state["r"]
        set_cell(ws, r, 1, "类别", font=Font(name=FONT, size=10, bold=True), fl=light, align=CENTER)
        set_cell(ws, r, 2, "作业单原文要求", font=Font(name=FONT, size=10, bold=True), fl=light, align=CENTER)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        set_cell(ws, r, 3, "本计划的安排", font=Font(name=FONT, size=10, bold=True), fl=light, align=CENTER)
        set_cell(ws, r, 4, None, fl=light)
        set_cell(ws, r, 5, None, fl=light)
        ws.row_dimensions[r].height = 18
        state["r"] = r + 1

    def item(sub, task, plan, light, dark, height=30):
        ws, r = state["ws"], state["r"]
        set_cell(ws, r, 1, sub, font=Font(name=FONT, size=10.5, bold=True, color="FFFFFF"), fl=dark, align=CENTER)
        set_cell(ws, r, 2, task, font=Font(name=FONT, size=11), fl=light, align=LEFT)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        set_cell(ws, r, 3, plan, font=Font(name=FONT, size=11), fl=light, align=LEFT)
        set_cell(ws, r, 4, None, fl=light)
        set_cell(ws, r, 5, None, fl=light)
        ws.row_dimensions[r].height = height
        state["r"] = r + 1

    def tip(text):
        ws, r = state["ws"], state["r"]
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        set_cell(ws, r, 1, text, font=Font(name=FONT, size=10, bold=True, color="7F6000"), fl=NOTE_BG, align=LEFT)
        for cc in range(2, 6):
            set_cell(ws, r, cc, None, fl=NOTE_BG)
        ws.row_dimensions[r].height = 30
        state["r"] = r + 1

    # ---------------- 总览1:语文 / 数学 ----------------
    start_sheet("总览1·语文数学",
                "四年级暑假作业总览(一)语文·数学 | 2026年7月13日—8月31日(共8周,9月1日开学)", first=True)

    section("一、语文暑假作业(完成时间自主安排;不含周六日和国家法定节假日)", CHN_D)
    head_row(CHN_L)
    item("基础1", "日积月累墨飘香:每日抄写1课《词语表》(语文书P142-144)中的词,每个词两遍,写在生字本上。要求:书写规范、工整、美观(必做题)",
         "每周一至周五每日1课,已排入每周计划页", CHN_L, CHN_D, 48)
    item("基础2", "吟诵经典入人心:背诵《208篇》的初中部分,不少于20篇。要求:声音洪亮,字音准确(必做题)",
         "第1~7周每周一、三、五各背1篇,共21篇(≥20篇)", CHN_L, CHN_D, 36)
    item("基础3", "练字闯关磨心性:每日自主练习一页字帖,并用“练字记录表”做好自我评价,开学后装订成册上交(个别同学选做;书写美观的同学免做)",
         "每周一至周五选做打卡", CHN_L, CHN_D, 54)
    item("实践1", "博览群书气自华:每日阅读自己感兴趣的书,并完成“我的阅读记录卡”(选做题,10分钟/每天)。推荐阅读:《中外民间故事》及其他名著、名家散文。开学实践活动:好书我推荐,2分钟,可制作PPT",
         "每周一至周五选做打卡;8月31日准备“好书我推荐”发言", CHN_L, CHN_D, 66)
    item("实践2", "写话提升勤耕笔(作文纸):①写写我的心爱之物:是什么、什么样、怎么得到、为什么成为,表达喜爱之情(必做)②慧眼看世界,妙笔游美景:写假期一处美景,按一定顺序写出景物动态变化(必做);有条件的同学录制旅行视频:横屏拍摄、声音清晰、不少于1分钟",
         "习作①第2周周五;习作②第5周周五", CHN_L, CHN_D, 92)
    item("实践3", "“心灵奇旅,观影推荐”:与家人乐享一场电影,自制“影片人物推荐卡”,内容包括人物名称、人物特点、人物画像及感受分享(选做题)。要求:注明电影名称,并用一件典型事例来介绍人物,充分说明他的特点",
         "第6周周五", CHN_L, CHN_D, 66)

    section("二、数学暑假作业", MATH_D)
    head_row(MATH_L)
    item("基础1", "口算(15~20道),题目自备;计算:三位数除以两位数的除法2道、小数乘法2道",
         "每周一至周五完成,已排入每周计划页", MATH_L, MATH_D)
    item("基础2", "每周完成1~2份综合练习卷,按8周计。三选一:①数练活页卷 ②《5.3全优卷》剩余试卷 ③可自备练习卷",
         "每周六1份必做(第1~6周共6份),第7周周六、周日各1份(第7、8份);每周日可加练1份(选做)", MATH_L, MATH_D, 34)
    item("选做", "挑选《教材全解》中喜欢的题目进行拓展练习;学有余力可与家长商定学习内容,自行安排拓展提高",
         "每周计划页设“选做”栏,学有余力时安排", MATH_L, MATH_D, 30)
    item("实践", "实践性作业任选两项,自己设计作业格式,配照片或表格,A4纸完成:①查找数学游戏,和朋友玩并记录规则 "
         "②调查小区占地面积、居住人口、活动面积,写调查报告 ③统计某一周(7天)气温,记录成表并绘统计图,提出数学问题并解答 "
         "④设计一幅美丽的密铺图案(单一图形或多图形组合)",
         "建议选③和④:第4周每天记气温、周日成图并提问解答;第6周周日设计密铺图案。也可换成①或②,时间照用", MATH_L, MATH_D, 62)
    tip("★ 语文、数学详细安排见“第N周”各工作表;完成一项在□打✓。总览(二)还有英语、体育和八周规划一览。")
    scale_rows(state["ws"], state["r"] - 1, target=1350)

    # ---------------- 总览2:英语 / 体育 / 八周一览 ----------------
    start_sheet("总览2·英语体育",
                "四年级暑假作业总览(二)英语·体育·八周一览 | 2026年7月13日—8月31日(共8周)")

    section("三、英语暑假作业(不含周六日和法定节假日;开学报到第一天上交学习自评单)", ENG_D)
    head_row(ENG_L)
    item("必做①", "每天大声朗读三下课本,共计24天,每次5分钟",
         "第1~6周,每周一至周四各1次(6周×4次=24次)", ENG_L, ENG_D)
    item("必做②", "每日抄写四上和四下单词或词组并造句,共计24天,每次10分钟",
         "与朗读同日进行:第1~6周,每周一至周四(24次)", ENG_L, ENG_D)
    item("必做③", "每日完成英语阅读题3篇,共计15天,每次15分钟",
         "第1~5周,每周一、三、五各1次(5周×3次=15次)", ENG_L, ENG_D)
    item("必做④", "学科实践活动:选择3~4首你喜爱的古诗,完成A3海报作品,共计5天,每次15分钟(古诗最好是自己学过的中文古诗)",
         "第7周周一至周五,每天15分钟分步完成:选诗→版面→抄写→配图→完善", ENG_L, ENG_D, 34)
    item("选做①", "学唱一首英文歌曲,歌曲自选,共计2天,每次5分钟",
         "第5周周五、第6周周五", ENG_L, ENG_D)
    item("选做②", "英语试听练习,一个英文动画电影或BBC记录片,内容自选,共计2天,每次不超过10分钟",
         "第7周周三、周五", ENG_L, ENG_D)
    item("自评单", "开学报到第一天上交《暑假英语学习自评单》:是否完成4项必做、是否完成2项选做,并写对下学期英语课的想法和建议",
         "8月31日填写,9月1日上交", ENG_L, ENG_D, 34)

    section("四、体育家庭任务单(第1~7周:7月13日—8月28日,周一至周五;做好记录,开学上交,学校评优体奖章)", PE_D)
    head_row(PE_L)
    item("周一", "①准备热身活动 ②一分钟计时跳绳×4组,并记录成绩 ③坐位体前屈拉伸(膝盖伸直、手触脚尖)1分钟",
         "已排入每周计划页“体育”栏", PE_L, PE_D)
    item("周二", "①准备热身活动 ②跳绳200个/组×4组 ③一分钟仰卧起坐2组,并记录成绩 ④坐位体前屈拉伸,膝盖伸直,手触脚尖1分钟",
         "同上", PE_L, PE_D)
    item("周三", "①准备热身运动 ②一分钟计时跳绳×4组,并记录成绩 ③肺活量练习(吸足气吹气球3组) ④放松拉伸",
         "同上", PE_L, PE_D)
    item("周四", "①准备热身活动 ②3分钟耐力跳绳(自选音乐),并记录跳绳总个数 ③放松拉伸",
         "同上", PE_L, PE_D)
    item("周五", "①准备热身活动 ②耐力跑2公里(避开中午炎热时段,注意防暑降温),并记录 ③放松拉伸",
         "同上", PE_L, PE_D)
    item("要求", "跳绳每组目标200个/分钟;仰卧起坐目标49个/分钟;坐位体前屈膝盖伸直手指超过脚尖;有场地条件可自行进行加速跑练习",
         "已写入每周计划页备注", PE_L, PE_D, 30)
    item("注明", "任务单末尾填写:跳绳最好的一次成绩、跑步公里数及所用时间、仰卧起坐最好的一次成绩;每周家长点评签字",
         "每周计划页有成绩记录行和家长签字行;第8周页统一核对填写", PE_L, PE_D, 30)

    section("五、八周规划一览(每周详细安排见对应工作表,打印各周页贴墙即可)", "7030A0")
    ws, r = state["ws"], state["r"]
    weeks_summary = [
        ("第1周 7.13–7.19", "抄词(一~五)/ 背诵3篇(一三五)/ 选做阅读、字帖", "口算计算(一~五)/ 试卷第1份(六)", "朗读+抄写造句(一~四)/ 阅读3篇(一三五)", "按课表锻炼+记成绩+家长签字"),
        ("第2周 7.20–7.26", "同上 + ★习作①心爱之物(五)", "口算计算 / 试卷第2份(六)", "同第1周", "同上"),
        ("第3周 7.27–8.2", "抄词 / 背诵3篇 / 选做阅读、字帖", "口算计算 / 试卷第3份(六)", "同第1周", "同上"),
        ("第4周 8.3–8.9", "同上", "口算计算 / 试卷第4份(六)/ ★实践每天记气温,周日绘统计图", "同第1周", "同上"),
        ("第5周 8.10–8.16", "同上 + ★习作②假期美景(五)", "口算计算 / 试卷第5份(六)", "朗读+抄写(一~四)/ 阅读最后一周 / ★选做英文歌(五)", "同上"),
        ("第6周 8.17–8.23", "同上 + ★选做影片推荐卡(五)", "口算计算 / 试卷第6份(六)/ ★实践周日设计密铺图案", "朗读+抄写最后一周 / ★英文歌(五)", "同上"),
        ("第7周 8.24–8.30", "抄词 / 背诵累计满21篇", "口算计算 / ★试卷第7份(六)、第8份(日)", "★A3古诗海报(一~五)/ ★选做试听(三、五)", "任务单最后一周,填“注明”最好成绩"),
        ("第8周 8.31", "核对背诵/习作,准备“好书我推荐”", "清点8份试卷+2项实践作业", "填写自评单(9月1日上交)", "自由锻炼,补齐签字"),
    ]
    set_cell(ws, r, 1, "周次", font=Font(name=FONT, size=10, bold=True, color="FFFFFF"), fl="7030A0", align=CENTER)
    set_cell(ws, r, 2, "语文", font=Font(name=FONT, size=10, bold=True, color="FFFFFF"), fl=CHN_D, align=CENTER)
    set_cell(ws, r, 3, "数学", font=Font(name=FONT, size=10, bold=True, color="FFFFFF"), fl=MATH_D, align=CENTER)
    set_cell(ws, r, 4, "英语", font=Font(name=FONT, size=10, bold=True, color="FFFFFF"), fl=ENG_D, align=CENTER)
    set_cell(ws, r, 5, "体育", font=Font(name=FONT, size=10, bold=True, color="FFFFFF"), fl=PE_D, align=CENTER)
    ws.row_dimensions[r].height = 20
    r += 1
    for i, (wk, c, m, e, p) in enumerate(weeks_summary):
        lt = "FFFFFF" if i % 2 == 0 else "F3F0F8"
        set_cell(ws, r, 1, wk, font=Font(name=FONT, size=10, bold=True), fl=PRAC_L, align=CENTER)
        set_cell(ws, r, 2, c, font=Font(name=FONT, size=10), fl=lt, align=LEFT)
        set_cell(ws, r, 3, m, font=Font(name=FONT, size=10), fl=lt, align=LEFT)
        set_cell(ws, r, 4, e, font=Font(name=FONT, size=10), fl=lt, align=LEFT)
        set_cell(ws, r, 5, p, font=Font(name=FONT, size=10), fl=lt, align=LEFT)
        ws.row_dimensions[r].height = 34
        r += 1
    state["r"] = r
    tip("★ 使用方法:每周日晚打印下一周的“第N周”工作表(A4横向,已设置好一页打印),贴在书桌前;"
        "孩子每完成一项在□打✓,周日晚家长检查并签字。语文、英语作业不含周六日及法定节假日(暑期内无法定节假日)。")
    scale_rows(state["ws"], state["r"] - 1, target=1350)


def main():
    wb = Workbook()
    build_overview(wb)
    for w in range(1, 8):
        build_week(wb, w)
    build_week8(wb)
    out = "四年级暑假作业计划表.xlsx"
    wb.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main()
